from pathlib import Path
from openpyxl import load_workbook
base = Path(r"C:\Users\HI MATEBOOK 14\Desktop\Dream AI 素材库\Dream AI 文档")
p = sorted(base.glob("*.xlsx"), key=lambda x: x.stat().st_mtime)[-1]
print(p.name)
wb = load_workbook(p)
for ws in wb.worksheets:
    print("sheet:", ws.title, ws.max_row, "x", ws.max_column)
    for row in ws.iter_rows(values_only=True):
        print([str(v) if v is not None else "" for v in row])