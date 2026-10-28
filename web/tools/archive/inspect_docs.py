from pathlib import Path
from itertools import islice
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

base = Path(r"C:\Users\HI MATEBOOK 14\Desktop\Dream AI 素材库\Dream AI 文档")
for p in sorted(base.iterdir()):
    print("\n==", p.name, p.stat().st_size)
    if p.suffix == ".docx":
        d = Document(p)
        print(" paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
        for para in d.paragraphs[:8]:
            if para.text.strip(): print("  ", para.text[:90])
        if d.tables:
            t = d.tables[0]
            print("  table:", len(t.rows), "x", len(t.columns), "|", " / ".join(c.text for c in t.rows[0].cells))
    elif p.suffix == ".xlsx":
        wb = load_workbook(p)
        for ws in wb.worksheets:
            print(" sheet:", ws.title, ws.max_row, "x", ws.max_column)
            for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True):
                print("   ", [str(v)[:18] if v else "" for v in row])
    elif p.suffix == ".pptx":
        prs = Presentation(p)
        print(" slides:", len(prs.slides))
        for slide in islice(prs.slides, 2):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.replace("\n", " / ")[:90])
            print("  ", " · ".join(texts)[:240])
    elif p.suffix == ".pdf":
        import fitz
        doc = fitz.open(p)
        print(" pages:", len(doc), "head:", p.read_bytes()[:5])
        print("  text:", doc[0].get_text()[:240].replace("\n", " / "))